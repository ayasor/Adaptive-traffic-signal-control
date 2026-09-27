"""Small helper layer over python-docx used by build_thesis.py."""

from __future__ import annotations

import hashlib
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

HERE = Path(__file__).resolve().parent
EQ_DIR = HERE / "equations"
ACCENT = RGBColor(0x0E, 0x5E, 0x6F)
MUTED = RGBColor(0x52, 0x51, 0x4E)
RUNNING_TITLE = "Urban traffic optimisation"

plt.rcParams["mathtext.fontset"] = "cm"


class Doc:
    def __init__(self) -> None:
        self.d = Document()
        self.fig_no = 0
        self.tab_no = 0
        self.eq_no = 0
        self.headings: list[tuple[int, str]] = []
        self.toc_pages: dict[str, int] = {}
        self._setup()

    # ------------------------------------------------------------------ setup
    def _setup(self) -> None:
        sec = self.d.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)
        st = self.d.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.2
        for lvl, size in ((1, 18), (2, 14), (3, 12), (4, 11)):
            h = st[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = ACCENT
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            h.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.size = Pt(9)
        cap.font.italic = True
        cap.font.color.rgb = MUTED
        cap.font.bold = False

    def running_header_and_page_numbers(self, section) -> None:
        section.different_first_page_header_footer = False
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = RUNNING_TITLE
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = MUTED
            r.italic = True
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._field(fp, "PAGE", size=9)

    # --------------------------------------------------------------- blocks
    def h(self, text: str, level: int = 1, page_break: bool = False):
        if page_break:
            self.page_break()
        if level <= 2 and text not in ("Contents",):
            self.headings.append((level, text))
        return self.d.add_heading(text, level=level)

    def p(self, text: str = "", *, italic=False, bold=False, align=None, size=None,
          color=None, space_after=None, style=None):
        """Paragraph; **bold** and *italic* inline markers are supported."""
        par = self.d.add_paragraph(style=style)
        self._inline(par, text, italic=italic, bold=bold, size=size, color=color)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify" or align is None:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        return par

    def bullets(self, items: list[str], numbered: bool = False) -> None:
        for n, it in enumerate(items, 1):
            if numbered:
                par = self.d.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.9)
                par.paragraph_format.first_line_indent = Cm(-0.6)
                self._inline(par, f"{n}.\u00a0\u00a0" + it)
            else:
                par = self.d.add_paragraph(style="List Bullet")
                self._inline(par, it)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.space_after = Pt(3)

    def note(self, title: str, text: str) -> None:
        """Shaded call-out box (one-cell table)."""
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.rows[0].cells[0]
        _shade(cell, "E8F1F2")
        _cell_borders(cell, "0E5E6F")
        cp = cell.paragraphs[0]
        r = cp.add_run(title)
        r.bold = True
        r.font.color.rgb = ACCENT
        for para in text.split("\n\n"):
            np_ = cell.add_paragraph()
            self._inline(np_, para)
            np_.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            np_.paragraph_format.space_after = Pt(4)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, path: Path, caption: str, width_cm: float = 15.0) -> None:
        self.fig_no += 1
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.d.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._inline(cap, f"Figure {self.fig_no}: {caption}")

    def equation(self, latex: str, numbered: bool = True, fontsize: int = 15) -> None:
        png = render_equation(latex, fontsize)
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(8)
        w_in = _png_width_in(png)
        par.add_run().add_picture(str(png), width=Cm(min(w_in * 2.54, 15.5)))
        if numbered:
            self.eq_no += 1
            r = par.add_run(f"\t({self.eq_no})")
            r.font.color.rgb = MUTED
            tabs = par.paragraph_format.tab_stops
            tabs.add_tab_stop(Cm(16))

    def table(self, header: list[str], rows: list[list], caption: str | None = None,
              widths_cm: list[float] | None = None, font_size: int = 9,
              align_numbers: bool = True) -> None:
        if caption:
            self.tab_no += 1
            cap = self.d.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
            self._inline(cap, f"Table {self.tab_no}: {caption}")
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for j, text in enumerate(header):
            c = t.rows[0].cells[j]
            _shade(c, "0E5E6F")
            c.paragraphs[0].text = ""
            r = c.paragraphs[0].add_run(str(text))
            r.bold = True
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for i, row in enumerate(rows, 1):
            for j, val in enumerate(row):
                c = t.rows[i].cells[j]
                if i % 2 == 0:
                    _shade(c, "F3F3F1")
                par = c.paragraphs[0]
                self._inline(par, str(val), size=font_size)
                if align_numbers and j > 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _repeat_header(t.rows[0])
        if widths_cm:
            t.autofit = False
            _grid(t, widths_cm)
            for row in t.rows:
                for j, w in enumerate(widths_cm):
                    row.cells[j].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self) -> None:
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def new_section(self):
        return self.d.add_section(WD_SECTION.NEW_PAGE)

    def toc(self, entries: list[tuple[int, str]], pages: dict[str, int]) -> None:
        """Static table of contents (page numbers from a previous render)."""
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        for level, text in entries:
            par = self.d.add_paragraph()
            pf = par.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0 if level > 1 else 3)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.6 * (level - 1))
            pf.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            r = par.add_run(f"{text}\t{pages.get(text, 0)}")
            r.bold = level == 1
            r.font.size = Pt(10.5 if level == 1 else 10)

    def save(self, path: Path) -> None:
        self.d.save(str(path))

    # ------------------------------------------------------------- internals
    def _inline(self, par, text, italic=False, bold=False, size=None, color=None) -> None:
        import re
        tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
        for tok in tokens:
            if not tok:
                continue
            b, i, mono = bold, italic, False
            if tok.startswith("**"):
                tok, b = tok[2:-2], True
            elif tok.startswith("*"):
                tok, i = tok[1:-1], True
            elif tok.startswith("`"):
                tok, mono = tok[1:-1], True
            r = par.add_run(tok)
            r.bold, r.italic = b, i
            if mono:
                r.font.name = "Consolas"
                r.font.size = Pt((size or 11) - 1)
            elif size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color

    def _field(self, par, instr: str, placeholder: str = "", size: int | None = None) -> None:
        run = par.add_run()
        f1 = OxmlElement("w:fldChar")
        f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        f2 = OxmlElement("w:fldChar")
        f2.set(qn("w:fldCharType"), "separate")
        run._r.append(f1)
        run._r.append(it)
        run._r.append(f2)
        r2 = par.add_run(placeholder or "1")
        if size:
            r2.font.size = Pt(size)
            run.font.size = Pt(size)
        r3 = par.add_run()
        f3 = OxmlElement("w:fldChar")
        f3.set(qn("w:fldCharType"), "end")
        r3._r.append(f3)


def render_equation(latex: str, fontsize: int = 15) -> Path:
    EQ_DIR.mkdir(exist_ok=True)
    name = hashlib.sha1(f"{latex}{fontsize}".encode()).hexdigest()[:12] + ".png"
    out = EQ_DIR / name
    if not out.exists():
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=fontsize)
        fig.savefig(out, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=False,
                    facecolor="white")
        plt.close(fig)
    return out


def _png_width_in(png: Path) -> float:
    from PIL import Image
    with Image.open(png) as im:
        return im.size[0] / 300


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _cell_borders(cell, color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def _grid(table, widths_cm: list[float]) -> None:
    tbl = table._tbl
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
        grid.append(gc)
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def _repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)
